# build_docs.py: turns the submission's Markdown files into Word files that
# import cleanly into Google Docs (headings, tables, pictures, Thai font).
#   python submission\build_docs.py
# Writes 2.1_Proposal\Proposal.docx and 2.3_Design_Document\Design_Document.docx.
# Only the Markdown this folder uses is understood: # headings, paragraphs,
# * / - / 1. lists, | tables |, ![caption](picture), > quotes, ``` code,
# **bold**, *italic*, `code`. Plus three directives (one per line):
#   <!-- CODE: path | first-line regex | last-line regex -->
#        lines of a source file (path from the project folder) with their
#        real line numbers: from the first line matching the first regex to
#        the next line matching the second one
#   <!-- SOURCE: path -->     a whole source file (the code appendix)
#   <!-- PAGEBREAK -->
import os, re, sys
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
FONT, MONO = 'Sarabun', 'Consolas'
ACCENT = RGBColor(0x4C, 0x1D, 0x95)

DOCS = [
    ('2.1_Proposal/proposal.md', '2.1_Proposal/Proposal.docx', False),
    ('2.3_Design_Document/design_document.md', '2.3_Design_Document/Design_Document.docx', True),
]


def set_font(run_or_style, name, size=None, bold=None, color=None):
    f = run_or_style.font
    f.name = name
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if color is not None:
        f.color.rgb = color
    el = run_or_style.element if hasattr(run_or_style, 'element') else run_or_style._element
    rpr = el.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    for k in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        fonts.set(qn(k), name)
    if size:
        szcs = rpr.find(qn('w:szCs'))
        if szcs is None:
            szcs = OxmlElement('w:szCs')
            rpr.append(szcs)
        szcs.set(qn('w:val'), str(int(size * 2)))
    if bold:
        if rpr.find(qn('w:bCs')) is None:
            rpr.append(OxmlElement('w:bCs'))


def setup(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.orientation = WD_ORIENT.PORTRAIT
    for side in ('left_margin', 'right_margin'):
        setattr(sec, side, Cm(2.2))
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
    st = doc.styles
    set_font(st['Normal'], FONT, 11)
    st['Normal'].paragraph_format.space_after = Pt(4)
    st['Normal'].paragraph_format.line_spacing = 1.15
    sizes = {1: 22, 2: 17, 3: 14, 4: 12}
    for lv, sz in sizes.items():
        s = st[f'Heading {lv}']
        set_font(s, FONT, sz, True, ACCENT if lv <= 2 else RGBColor(0x1D, 0x24, 0x33))
        s.paragraph_format.space_before = Pt(14 if lv <= 2 else 10)
        s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.keep_with_next = True
    for name in ('List Bullet', 'List Number'):
        set_font(st[name], FONT, 11)


INLINE = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)')


def add_inline(p, text, size=None, bold=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2]); set_font(r, FONT, size, True)
        elif part.startswith('`') and part.endswith('`'):
            r = p.add_run(part[1:-1]); set_font(r, MONO, (size or 11) - 1, bold or None, RGBColor(0x9D, 0x17, 0x4D))
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            r = p.add_run(part[1:-1]); set_font(r, FONT, size, bold or None); r.italic = True
        else:
            r = p.add_run(part); set_font(r, FONT, size, bold or None)


def shade(cell, hexcol):
    tcpr = cell._element.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexcol)
    tcpr.append(sh)


def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip('|').split('|')] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r':?-{2,}:?', c) or c == '' for c in r)] if len(cells) > 1 else cells
    hdr, body = cells[0], cells[1:]
    empty_hdr = all(c == '' for c in hdr)
    data = body if empty_hdr else [hdr] + body
    ncol = max(len(r) for r in data)
    t = doc.add_table(rows=len(data), cols=ncol)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, r in enumerate(data):
        for j in range(ncol):
            c = t.cell(i, j)
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            add_inline(c.paragraphs[0], r[j] if j < len(r) else '', size=10, bold=(i == 0 and not empty_hdr))
            if i == 0 and not empty_hdr:
                shade(c, 'EDE9FE')
    # column widths in proportion to the longest text in each column
    weight = [max(min(max(len(r[j]) if j < len(r) else 0 for r in data), 45), 14) + 6 for j in range(ncol)]
    t.autofit = False
    for j in range(ncol):
        wcm = Cm(16.6 * weight[j] / sum(weight))
        for i in range(len(data)):
            t.cell(i, j).width = wcm
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_picture(doc, md_dir, cap, path):
    full = os.path.normpath(os.path.join(md_dir, path))
    if not os.path.exists(full):
        p = doc.add_paragraph(); add_inline(p, f'[รูป: {cap}]'); return
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    from PIL import Image
    w, h = Image.open(full).size
    width = 16.5 if w / h > 1.1 else 11.0
    if width * h / w > 21.5:                        # too tall for the page
        width = 21.5 * w / h
    p.add_run().add_picture(full, width=Cm(width))
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(cap); set_font(r, FONT, 10, None, RGBColor(0x5B, 0x64, 0x75)); r.italic = True


def code_lines(doc, lines, first_no=None, size=8.5):
    """source lines on a light grey background, Consolas, with line numbers"""
    for k, ln in enumerate(lines):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = 1.0
        pf.left_indent = Cm(0.2)
        ppr = p._element.get_or_add_pPr()
        sh = OxmlElement('w:shd')
        sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), 'F4F5F7')
        ppr.append(sh)
        if first_no is not None:
            r = p.add_run(f'{first_no + k:4d}  '); set_font(r, MONO, size, None, RGBColor(0x8A, 0x93, 0xA3))
        r = p.add_run(ln.rstrip('\n').replace('\t', '    ') or ' '); set_font(r, MONO, size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def code_excerpt(doc, spec):
    path, a, b = [x.strip() for x in spec.split('|')]
    src = open(os.path.join(ROOT, path), encoding='utf-8').read().split('\n')
    i = next(k for k, l in enumerate(src) if re.search(a, l))
    j = next(k for k in range(i, len(src)) if re.search(b, src[k]))
    cap = doc.add_paragraph()
    cap.paragraph_format.space_after = Pt(1); cap.paragraph_format.keep_with_next = True
    r = cap.add_run(f'{path}  บรรทัด {i + 1}-{j + 1}'); set_font(r, MONO, 8.5, True, RGBColor(0x4C, 0x1D, 0x95))
    code_lines(doc, src[i:j + 1], i + 1)


def source_file(doc, path):
    src = open(os.path.join(ROOT, path), encoding='utf-8').read().rstrip('\n').split('\n')
    h = doc.add_heading('', level=3)
    add_inline(h, f'`{path}` ({len(src):,} บรรทัด)')
    code_lines(doc, src, 1, size=7.5)


def new_numbering(doc, style_name):
    """a new w:num for a numbered list, so that each list starts again at 1"""
    numbering = doc.part.numbering_part.element
    numpr = doc.styles[style_name].element.pPr.find(qn('w:numPr'))
    num_id = numpr.find(qn('w:numId')).get(qn('w:val'))
    abstract = next(n.find(qn('w:abstractNumId')).get(qn('w:val'))
                    for n in numbering.findall(qn('w:num')) if n.get(qn('w:numId')) == num_id)
    new_id = max(int(n.get(qn('w:numId'))) for n in numbering.findall(qn('w:num'))) + 1
    num = OxmlElement('w:num'); num.set(qn('w:numId'), str(new_id))
    an = OxmlElement('w:abstractNumId'); an.set(qn('w:val'), abstract); num.append(an)
    ov = OxmlElement('w:lvlOverride'); ov.set(qn('w:ilvl'), '0')
    so = OxmlElement('w:startOverride'); so.set(qn('w:val'), '1'); ov.append(so); num.append(ov)
    numbering.append(num)
    return new_id


def set_num(p, num_id):
    ppr = p._element.get_or_add_pPr()
    numpr = OxmlElement('w:numPr')
    il = OxmlElement('w:ilvl'); il.set(qn('w:val'), '0'); numpr.append(il)
    ni = OxmlElement('w:numId'); ni.set(qn('w:val'), str(num_id)); numpr.append(ni)
    ppr.append(numpr)


def convert(md_path, out_path, chapter_breaks):
    md_dir = os.path.dirname(md_path)
    lines = open(md_path, encoding='utf-8').read().split('\n')
    if lines and lines[0].strip() == '---':
        end = lines.index('---', 1)
        lines = lines[end + 1:]
    doc = Document()
    setup(doc)
    i, para, first_h2 = 0, [], True
    list_num = None                                 # numbering of the numbered list in progress

    def flush():
        nonlocal para
        if para:
            p = doc.add_paragraph()
            add_inline(p, ' '.join(s.strip() for s in para))
            para = []

    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if s and not re.match(r'^\d+\.\s', ln) and not ln.startswith('  '):
            list_num = None                         # any other block ends a numbered list
        m = re.match(r'^<!--\s*(CODE|SOURCE|PAGEBREAK)\s*:?\s*(.*?)\s*-->$', s)
        if m:
            flush()
            if m.group(1) == 'CODE':
                code_excerpt(doc, m.group(2))
            elif m.group(1) == 'SOURCE':
                source_file(doc, m.group(2))
            else:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            i += 1; continue
        if s.startswith('```'):
            flush(); i += 1; blk = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                blk.append(lines[i]); i += 1
            code_lines(doc, blk, None, size=9)
            i += 1; continue
        m = re.match(r'^(#{1,4})\s+(.*)$', s)
        if m:
            flush()
            lv = len(m.group(1))
            if lv == 2 and chapter_breaks:
                if not first_h2:
                    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                elif first_h2:
                    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                first_h2 = False
            h = doc.add_heading('', level=lv)
            add_inline(h, m.group(2))
            if lv == 1:
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1; continue
        if s.startswith('|'):
            flush(); rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append(lines[i]); i += 1
            add_table(doc, rows); continue
        m = re.match(r'^!\[(.*?)\]\((.*?)\)$', s)
        if m:
            flush(); add_picture(doc, md_dir, m.group(1), m.group(2)); i += 1; continue
        m = re.match(r'^(\s*)([*-]|\d+\.)\s+(\[[ x]\]\s+)?(.*)$', ln)
        if m and s:
            flush()
            item = m.group(4)
            # continuation lines (indented, not a new item)
            while i + 1 < len(lines) and lines[i + 1].startswith('  ') and lines[i + 1].strip() and not re.match(r'^\s*([*-]|\d+\.)\s', lines[i + 1]):
                i += 1; item += ' ' + lines[i].strip()
            style = 'List Number' if m.group(2)[0].isdigit() else 'List Bullet'
            p = doc.add_paragraph(style=style)
            if style == 'List Number' and not m.group(1):
                if list_num is None:
                    list_num = new_numbering(doc, 'List Number')
                set_num(p, list_num)
            if m.group(1):
                p.paragraph_format.left_indent = Cm(1.6)
            box = '☐ ' if m.group(3) else ''
            add_inline(p, box + item)
            p.paragraph_format.space_after = Pt(2)
            i += 1; continue
        if s.startswith('>'):
            flush()
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.8)
            add_inline(p, s.lstrip('> ').strip()); i += 1; continue
        if s == '' or s == '---':
            flush(); i += 1; continue
        para.append(s); i += 1
    flush()
    doc.save(out_path)
    print('wrote', out_path)


if __name__ == '__main__':
    for md, out, br in DOCS:
        convert(os.path.join(HERE, md), os.path.join(HERE, out), br)
